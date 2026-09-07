"""Parse official GLD, IAU and SLV sponsor downloads without fixed offsets."""
import csv
import io
import json
import logging
import math
import re
import xml.etree.ElementTree as ET
from datetime import date, datetime, timedelta, timezone
from html.parser import HTMLParser

from collectors.base import HTTPCollector

logger = logging.getLogger(__name__)

FUNDS = {
    "GLD": ("GOLD", "https://api.spdrgoldshares.com/api/v1/historical-archive?exchange=NYSE&lang=en&product=gld"),
    "IAU": ("GOLD", "https://www.blackrock.com/varnish-api/blk-one01-product-data/product-data/api/v1/get-fund-document?appSubType=ISHARES&appType=PRODUCT_PAGE&component=fundDownload&locale=en_US&portfolioId=239561&targetSite=us-ishares&userType=individual"),
    "SLV": ("SILVER", "https://www.blackrock.com/varnish-api/blk-one01-product-data/product-data/api/v1/get-fund-document?appSubType=ISHARES&appType=PRODUCT_PAGE&component=fundDownload&locale=en_US&portfolioId=239855&targetSite=us-ishares&userType=individual"),
}
FALLBACKS = {
    "IAU": ("https://www.ishares.com/us/products/239561/ishares-gold-trust-fund",),
    "SLV": ("https://www.ishares.com/us/products/239855/ishares-silver-trust-fund",),
}
DATE_FIELDS = {"date", "effective date", "as of", "as of date", "fund holdings as of", "holdings as of"}


class _SponsorHTML(HTMLParser):
    """Keep visible leaf text and element boundaries for dated summary cards."""
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = {"tag": "root", "children": []}
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        node = {"tag": tag, "children": [], "parent": self.stack[-1]}
        self.stack[-1]["children"].append(node)
        if tag not in {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "track", "wbr"}:
            self.stack.append(node)

    def handle_endtag(self, tag):
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i]["tag"] == tag:
                del self.stack[i:]
                break

    def handle_data(self, data):
        text = " ".join(data.split())
        if text and not any(n["tag"] in {"script", "style", "noscript"} for n in self.stack):
            self.stack[-1]["children"].append(text)

    @staticmethod
    def tokens(node):
        if isinstance(node, str):
            return [node]
        return [part for child in node["children"] for part in _SponsorHTML.tokens(child)]

    def nodes(self, node=None):
        node = node or self.root
        for child in node["children"]:
            if isinstance(child, dict):
                yield from self.nodes(child)
        yield node


def _norm(value):
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(value).lower()).split())


def _num(value):
    """Parse sponsor-formatted numbers, while never converting missing data to zero."""
    if value is None:
        return None
    text = str(value).strip().strip('"').replace("\u00a0", " ")
    if text.lower() in ("", "-", "--", "n/a", "na", "not available"):
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.removeprefix("(").removesuffix(")").strip()
    # Only known decoration may surround the number.  This prevents accidentally
    # interpreting dates, identifiers, or arbitrary prose as holdings.
    match = re.fullmatch(r"(?:[$£€]\s*)?([+-]?[\d,]+(?:\.\d+)?)\s*(?:%|oz|ounces?|tonnes?|tons?)?", text, re.I)
    if not match:
        return None
    try:
        number = float(match.group(1).replace(",", ""))
        return (-number if negative else number) if math.isfinite(number) else None
    except ValueError:
        return None


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    value = str(value or "").strip().strip('"')
    value = re.sub(r"^(?:as\s+of|holdings\s+as\s+of)\s+", "", value, flags=re.I)
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:Z|[+-]\d{2}:\d{2})?", value):
        try:
            return datetime.fromisoformat(value).date()
        except ValueError:
            return None
    for fmt in ("%d-%b-%Y", "%d-%b-%y", "%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%b %d, %Y", "%B %d, %Y", "%b %d %Y", "%B %d %Y",
                "%m/%d/%Y", "%m/%d/%y", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    return None


def _measure(value, kind):
    """A declared field unit must agree with any unit printed by the sponsor."""
    text = str(value if value is not None else "").strip().strip('"').replace("\u00a0", " ")
    numeric = r"[+-]?[\d,]+(?:\.\d+)?"
    patterns = {
        "COUNT": numeric,
        "OUNCES": numeric + r"\s*(?:oz|(?:troy\s+)?ounces?)?",
        "TONNES": numeric + r"\s*(?:(?:metric\s+)?tonnes?)?",
        "USD": r"(?:(?:US\$|USD|\$)\s*)?" + numeric + r"\s*(?:USD)?",
    }
    if not re.fullmatch(patterns[kind], text, re.I):
        return None
    # _num understands $/oz/tonnes; strip the explicitly validated USD variants.
    text = re.sub(r"^(?:US\$|USD)\s*|\s*USD$", "", text, flags=re.I)
    text = re.sub(r"\s+(?:troy|metric)\s+", " ", text, flags=re.I)
    return _num(text)


def _matching(values, aliases, kind=None):
    """Pick an unambiguous semantic field; aliases are exact normalized names."""
    normalized_aliases = {_norm(alias) for alias in aliases}
    for key, value in values.items():
        if _norm(key) in normalized_aliases:
            parsed = _measure(value, kind) if kind else _num(value)
            if parsed is not None:
                return parsed
    return None


def _pairs(rows):
    """Return metadata from both ``key,value`` and ``header/value`` CSV blocks."""
    result = {}
    for row in rows:
        cells = [cell.strip() for cell in row]
        if len(cells) >= 2 and cells[0] and _norm(cells[1]) not in DATE_FIELDS | {"shares outstanding", "net assets", "ounces in trust", "tonnes in trust"}:
            result.setdefault(cells[0], cells[1])
    # Current iShares downloads contain a conventional field-name row followed
    # by its values in parts of the preamble.  Do not confuse the holdings table
    # with metadata: only recognised summary labels are promoted.
    summary = {"shares outstanding", "net assets", "net assets of fund",
               "ounces in trust", "tonnes in trust", "fund holdings as of",
               "holdings as of", "as of"}
    for header, values in zip(rows, rows[1:]):
        normalized = [_norm(cell) for cell in header]
        if len(values) == len(header) and sum(name in summary for name in normalized) >= 2:
            for key, value in zip(header, values):
                if _norm(key) in summary and str(value).strip():
                    result[str(key).strip()] = str(value).strip()
    return result


class MetalsETFCollector(HTTPCollector):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.diagnostics: dict[str, dict] = {}

    @staticmethod
    def _decode(response) -> str:
        raw = response.content
        # Current sponsor downloads are workbooks. Sniff their bytes, because
        # BlackRock also serves SpreadsheetML under the legacy Excel MIME type.
        if raw.startswith(b"PK\x03\x04"):
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            output = io.StringIO()
            writer = csv.writer(output)
            try:
                for sheet in workbook:
                    writer.writerows(sheet.iter_rows(values_only=True))
                    writer.writerow(["__SPONSOR_WORKSHEET_BOUNDARY__"])
            finally:
                workbook.close()
            return output.getvalue()
        if raw.startswith(b"\xd0\xcf\x11\xe0"):
            import xlrd
            workbook = xlrd.open_workbook(file_contents=raw)
            output = io.StringIO()
            writer = csv.writer(output)
            for sheet in workbook.sheets():
                for index in range(sheet.nrows):
                    writer.writerow([xlrd.xldate.xldate_as_datetime(cell.value, workbook.datemode)
                                     if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                                     for cell in sheet.row(index)])
                writer.writerow(["__SPONSOR_WORKSHEET_BOUNDARY__"])
            return output.getvalue()
        decoded = None
        for encoding in ("utf-8-sig", "utf-16", "latin-1"):
            try:
                decoded = raw.decode(encoding)
                break
            except (UnicodeDecodeError, UnicodeError):
                continue
        decoded = decoded if decoded is not None else response.text
        if "urn:schemas-microsoft-com:office:spreadsheet" in decoded[:10000]:
            namespace = "{urn:schemas-microsoft-com:office:spreadsheet}"
            root = ET.fromstring(decoded)
            output = io.StringIO()
            writer = csv.writer(output)
            for worksheet in root.iter(namespace + "Worksheet"):
                for row in worksheet.iter(namespace + "Row"):
                    values = []
                    for cell in row.findall(namespace + "Cell"):
                        index = int(cell.get(namespace + "Index", len(values) + 1))
                        values.extend([""] * max(0, index - len(values) - 1))
                        data = cell.find(namespace + "Data")
                        values.append(data.text if data is not None and data.text else "")
                    writer.writerow(values)
                writer.writerow(["__SPONSOR_WORKSHEET_BOUNDARY__"])
            return output.getvalue()
        if raw.startswith(b"%PDF"):
            raise ValueError("Sponsor returned a PDF, not fund summary data")
        return decoded

    def _diagnostic_record(self, fund, asset, source, fetched, status, error):
        # ``date`` is the storage/upsert partition, not an asserted sponsor date.
        # effective_date intentionally remains NULL when the sponsor did not supply one.
        return {"asset": asset, "fund": fund, "date": fetched.date(), "effective_date": None,
                "shares_outstanding": None, "ounces": None, "tonnes": None, "nav": None,
                "physical_holdings": None, "holdings_unit": None, "net_assets": None,
                "flow": None, "flow_status": "UNAVAILABLE", "source": source,
                "fetched_at": fetched, "published_at": None, "status": status,
                "error": str(error)[:500]}

    @staticmethod
    def _shape(content):
        """Build a bounded, non-sensitive schema description for field logs."""
        if re.search(r"<!doctype\s+html|<html", content[:1000], re.I):
            return {"type": "HTML", "row_count": 0}
        try:
            value = json.loads(content)
        except (json.JSONDecodeError, TypeError):
            rows = list(csv.reader(io.StringIO(content.lstrip("\ufeff").replace("\x00", ""))))
            keywords = ("shares outstanding", "net assets", "ounces", "tonnes",
                        "fund holdings", "holdings", "as of")
            candidates = []
            table_headers = []
            worksheet = 1
            for number, row in enumerate(rows, 1):
                if row == ["__SPONSOR_WORKSHEET_BOUNDARY__"]:
                    worksheet += 1
                    continue
                normalized_row = _norm(" ".join(row))
                normalized_cells = {_norm(cell) for cell in row}
                # Search every sheet and all columns for the actual data header;
                # a long legal/preamble sheet must not consume the log budget.
                # Only header names enter this bounded output, never data rows.
                if (normalized_cells & DATE_FIELDS and any(word in normalized_row for word in keywords)
                        and len(table_headers) < 12):
                    names = [_norm(cell)[:140] for cell in row[:32]
                             if re.search(r"[A-Za-z]", str(cell)) and _num(cell) is None and _date(cell) is None]
                    table_headers.append({"worksheet": worksheet, "line": number,
                                          "columns": len(row), "fields": names})
                # Schema only: values (including dollar amounts and dates) never
                # enter diagnostics. In key/value rows only the key is a field.
                if number <= 40 and len(candidates) < 15:
                    possible_headers = row[:1] if len(row) == 2 else row[:6]
                    names = [_norm(cell)[:60] for cell in possible_headers
                             if re.search(r"[A-Za-z]", str(cell)) and _num(cell) is None and _date(cell) is None]
                    if number <= 5 or any(word in normalized_row for word in keywords):
                        candidates.append({"line": number, "columns": len(row), "fields": names})
            return {"type": "CSV", "row_count": len(rows), "candidate_rows": candidates,
                    "table_headers": table_headers}
        items = value if isinstance(value, list) else next(
            (v for v in value.values() if isinstance(v, list)), []) if isinstance(value, dict) else []
        top = sorted(map(str, value.keys()))[:30] if isinstance(value, dict) else []
        keys = lambda item: sorted(map(str, item.keys()))[:30] if isinstance(item, dict) else []
        all_keys = {_norm(k) for item in items[:100] if isinstance(item, dict) for k in item}
        return {"type": "JSON", "top_level_keys": top, "row_count": len(items),
                "first_item_keys": keys(items[0]) if items else [],
                "last_item_keys": keys(items[-1]) if items else [],
                "weight_fields": sorted(k for k in all_keys if "weight" in k or "ounce" in k or "tonne" in k),
                "date_fields": sorted(k for k in all_keys if "date" in k)}

    def fetch_history(self, start_date: date, end_date: date) -> list[dict]:
        records = []
        self.diagnostics = {}
        if start_date > end_date:
            return records
        for fund, (asset, primary) in FUNDS.items():
            fetched = datetime.now(timezone.utc)
            diag = {"http_status": None, "raw_count": 0, "parsed_count": 0,
                    "normalized_count": 0, "status": "HTTP_ERROR", "attempts": []}
            for url in (primary, *FALLBACKS.get(fund, ())):
                attempt = {"source": url, "http_status": None, "status": "HTTP_ERROR",
                           "raw_count": 0, "parsed_count": 0, "normalized_count": 0, "error": None}
                try:
                    response = self.client.get(url, follow_redirects=True, headers={"User-Agent": "Mozilla/5.0"})
                    attempt["http_status"] = response.status_code
                    response.raise_for_status()
                    attempt["status"] = "HTTP_OK_SCHEMA_UNEXPECTED"
                    content = self._decode(response)
                    shape = self._shape(content)
                    attempt["schema"] = shape
                    attempt["raw_count"] = shape["row_count"]
                    logger.info("%s response schema=%s", fund, shape)
                    parsed = self._parse(fund, asset, str(response.url), content, start_date, end_date, fetched=fetched)
                    attempt["parsed_count"] = attempt["normalized_count"] = len(parsed)
                    attempt["field_counts"] = {field: sum(record.get(field) is not None for record in parsed)
                                               for field in ("physical_holdings", "ounces", "tonnes", "shares_outstanding", "net_assets")}
                    logger.info("%s parsed field coverage=%s", fund, attempt["field_counts"])
                    if parsed:
                        attempt["status"] = "HTTP_OK_PARSE_OK"
                        records.extend(parsed)
                    else:
                        normalized_content = _norm(content[:10000])
                        known_schema = any(marker in normalized_content for marker in
                                           ("shares outstanding", "fund holdings as of", "total gold in trust", "ounces in trust"))
                        attempt["status"] = "HTTP_OK_PARSE_EMPTY" if known_schema else "HTTP_OK_SCHEMA_UNEXPECTED"
                except Exception as exc:
                    attempt["error"] = type(exc).__name__
                diag["attempts"].append(attempt)
                diag.update(attempt)
                if attempt["status"] == "HTTP_OK_PARSE_OK":
                    break
            self.diagnostics[fund] = diag
            logger.info("%s HTTP=%s raw rows=%d parsed rows=%d normalized records=%d status=%s",
                        fund, diag["http_status"], diag["raw_count"], diag["parsed_count"],
                        diag["normalized_count"], diag["status"])
        return records

    def _record(self, fund, asset, source, day, fetched, values):
        shares = _matching(values, ("shares outstanding", "total shares outstanding", "shares outstanding as of", "shares outstanding end of day"), "COUNT")
        ounces = _matching(values, ("total gold in trust in ounces", "total silver in trust in ounces",
                           "total gold in trust (ounces)", "total silver in trust (ounces)",
                           "total net asset value ounces in the trust",
                           "total ounces", "ounces in trust", "ounces", "fine ounces"), "OUNCES")
        tonnes = _matching(values, ("total gold in trust in tonnes", "total silver in trust in tonnes",
                           "total gold in trust (tonnes)", "total silver in trust (tonnes)",
                           "total net asset value tonnes in the trust",
                           "tonnes in trust", "tonnes", "metric tonnes"), "TONNES")
        nav = _matching(values, ("total net asset value in the trust", "total net asset value", "total net assets", "net assets",
                        "net assets of fund", "market value"), "USD")
        shares, ounces, tonnes, nav = [v if v is not None and v >= 0 else None
                                       for v in (shares, ounces, tonnes, nav)]
        if all(v is None for v in (shares, ounces, tonnes, nav)):
            return None
        physical, unit = (ounces, "OUNCES") if ounces is not None else ((tonnes, "TONNES") if tonnes is not None else (None, None))
        published_value = next((v for k, v in values.items() if _norm(k) in ("published at", "publication date")), None)
        published_day = _date(published_value)
        return {"asset": asset, "fund": fund, "date": day or fetched.date(), "effective_date": day,
                "shares_outstanding": shares, "ounces": ounces, "tonnes": tonnes, "nav": nav,
                "physical_holdings": physical, "holdings_unit": unit, "net_assets": nav,
                "flow": None, "flow_status": "UNAVAILABLE", "source": source, "fetched_at": fetched,
                "published_at": datetime.combine(published_day, datetime.min.time(), tzinfo=timezone.utc) if published_day else None,
                "status": "OK", "error": None}

    def _parse(self, fund, asset, source, content, start_date, end_date, fetched=None):
        fetched = fetched or datetime.now(timezone.utc)
        if start_date > end_date:
            return []
        if "__SPONSOR_WORKSHEET_BOUNDARY__" in content:
            # Metadata dates and values belong to one worksheet. In particular,
            # do not let a later worksheet's ounces inherit the first one's date.
            result = []
            for section in re.split(r"(?m)^__SPONSOR_WORKSHEET_BOUNDARY__\r?$", content):
                if section.strip():
                    result.extend(self._parse(fund, asset, source, section, start_date, end_date, fetched))
            merged = {}
            for index, record in enumerate(result):
                key = record["effective_date"] or ("undated", index)
                if key not in merged:
                    merged[key] = record
                else:
                    merged[key].update({k: v for k, v in record.items() if v is not None})
            return sorted(merged.values(), key=lambda row: row["date"])
        if re.search(r"<!doctype\s+html|<html", content[:1000], re.I):
            return self._parse_html_summary(fund, asset, source, content, start_date, end_date, fetched)
        try:
            payload = json.loads(content)
        except (json.JSONDecodeError, TypeError):
            payload = None
        if payload is not None:
            records = self._parse_json_summaries(fund, asset, source, payload, fetched)
            return [record for record in records if record["effective_date"] is None or
                    start_date <= record["effective_date"] <= end_date]
        rows = list(csv.reader(io.StringIO(content.lstrip("\ufeff").replace("\x00", ""))))
        result = []

        # GLD archive: detect the semantic header rather than assuming its line number.
        headers = None
        date_key = None
        archive = False
        for row in rows:
            if row == ["__SPONSOR_WORKSHEET_BOUNDARY__"]:
                headers = date_key = None
                continue
            normalized = {_norm(c) for c in row}
            # A metadata key/date pair is not a table header. Date and at least
            # one recognised measured summary field must share the header row.
            table_dates = DATE_FIELDS - {"fund holdings as of", "holdings as of"}
            if normalized & table_dates and normalized & {
                "shares outstanding", "total shares outstanding", "shares outstanding as of",
                "total gold in trust in ounces", "total gold in trust in tonnes",
                "total gold in trust ounces", "total gold in trust tonnes",
                "total silver in trust ounces", "total silver in trust tonnes",
                "total net asset value ounces in the trust", "total net asset value tonnes in the trust",
                "total net asset value in the trust", "total net assets", "net assets",
                "net assets of fund", "ounces in trust", "tonnes in trust",
            }:
                headers = [c.strip() for c in row]
                date_key = next(h for h in headers if _norm(h) in table_dates)
                archive = True
                continue
            if headers:
                values = dict(zip(headers, row))
                day = _date(values.get(date_key))
                if day and start_date <= day <= end_date:
                    record = self._record(fund, asset, source, day, fetched, values)
                    if record:
                        result.append(record)
        if result:
            # Workbooks may repeat the same day's summary across worksheets.
            merged = {}
            for record in result:
                day = record["effective_date"]
                if day not in merged:
                    merged[day] = record
                else:
                    merged[day].update({k: v for k, v in record.items() if v is not None})
            return [merged[day] for day in sorted(merged)]
        if archive:
            return []

        # iShares has metadata before its actual holdings header. Both can move.
        pairs = _pairs(rows)
        asof_value = next((v for k, v in pairs.items() if _norm(k) in
                          ("fund holdings as of", "holdings as of", "as of", "as of date", "effective date")), None)
        day = _date(asof_value)
        if day and not start_date <= day <= end_date:
            return []
        values = dict(pairs)
        holdings_header = next((i for i, row in enumerate(rows)
                                if "quantity" in {_norm(c) for c in row} and
                                ({"name", "asset class", "ticker"} & {_norm(c) for c in row})), None)
        if holdings_header is not None:
            headers = [_norm(c) for c in rows[holdings_header]]
            candidates = []
            for row in rows[holdings_header + 1:]:
                item = dict(zip(headers, row))
                identity = " ".join((item.get("ticker", ""), item.get("name", ""), item.get("asset class", ""))).lower()
                metal = "gold" if asset == "GOLD" else "silver"
                if metal in identity or item.get("ticker", "").upper() in ({"XAU", "GOLD"} if asset == "GOLD" else {"XAG", "SLV"}):
                    quantity = _measure(item.get("quantity"), "OUNCES")
                    if quantity is not None:
                        candidates.append((quantity, item))
            # Never sum a bar list: only a single commodity holding is unambiguous.
            if len(candidates) == 1:
                quantity, item = candidates[0]
                values["ounces"] = quantity
                market_value = _measure(item.get("market value"), "USD")
                if market_value is not None and _matching(values, ("total net assets", "net assets", "net assets of fund")) is None:
                    values["net assets"] = market_value
        record = self._record(fund, asset, source, day, fetched, values)
        return [record] if record else []

    def _parse_json_summaries(self, fund, asset, source, payload, fetched):
        """Parse an explicit sponsor summary; deliberately never sum bar rows."""
        objects = []
        def visit(value, depth=0):
            if depth > 4:
                return
            if isinstance(value, dict):
                objects.append(value)
                for child in value.values():
                    visit(child, depth + 1)
            elif isinstance(value, list):
                for child in value:
                    visit(child, depth + 1)
        visit(payload)
        records = []
        for values in objects:
            normalized = {_norm(key): value for key, value in values.items()}
            # A bar's weight or market value is never a fund-level total.
            if not set(normalized) & {"shares outstanding", "total shares outstanding",
                    "ounces in trust", "tonnes in trust", "total gold in trust in ounces",
                    "total gold in trust in tonnes", "net assets", "net assets of fund",
                    "total net assets", "total net asset value in the trust"}:
                continue
            day_value = next((value for key, value in normalized.items()
                              if key in DATE_FIELDS), None)
            day = _date(day_value)
            record = self._record(fund, asset, source, day, fetched, values)
            if record:
                records.append(record)
        return records

    def _parse_html_summary(self, fund, asset, source, content, start_date, end_date, fetched):
        """Read explicitly dated sponsor fact cards without borrowing page dates."""
        parser = _SponsorHTML()
        parser.feed(content)
        expected_name = {"IAU": "ishares gold trust", "SLV": "ishares silver trust", "GLD": "spdr gold"}[fund]
        if expected_name not in _norm(" ".join(parser.tokens(parser.root))):
            return []
        aliases = {"shares outstanding", "total shares outstanding", "net assets of fund",
                   "total net asset value in the trust", "ounces in trust", "tonnes in trust"}
        summaries = {}
        for label_node in parser.nodes():
            label_tokens = parser.tokens(label_node)
            if len(label_tokens) != 1 or _norm(label_tokens[0]) not in aliases:
                continue
            label = _norm(label_tokens[0])
            node = label_node.get("parent")
            numbers = []
            # The first local container holding this label and its value is its
            # fact card. A missing date here cannot be supplied by a page footer.
            while node and node["tag"] not in {"root", "body", "html", "footer", "main"}:
                tokens = parser.tokens(node)
                labels = {_norm(token) for token in tokens} & aliases
                if len(tokens) > 16 or len(labels) != 1:
                    break
                numbers = [token for token in tokens if _date(token) is None and _num(token) is not None]
                if numbers:
                    break
                node = node.get("parent")
            if not node or len(numbers) != 1:
                continue
            text = " ".join(tokens)
            # Only a date introduced by the sponsor's 'as of' label qualifies.
            date_match = re.search(r"\bas\s+of\s+([A-Za-z]+\s+\d{1,2},?\s+\d{4}|\d{4}-\d{2}-\d{2}|\d{1,2}/\d{1,2}/\d{4})\b", text, re.I)
            day = _date(date_match.group(1)) if date_match else None
            if day is None or not start_date <= day <= end_date:
                continue
            decoration = r"(?:[$£€%]|[A-Z]{3}|oz|(?:troy\s+)?ounces?|(?:metric\s+)?tonnes?|tons?)"
            raw_value = " ".join(token for token in tokens if token == numbers[0]
                                 or re.fullmatch(decoration, token))
            summaries.setdefault(day, {}).setdefault(label, raw_value)
        records = [self._record(fund, asset, source, day, fetched, values)
                   for day, values in sorted(summaries.items())]
        return [record for record in records if record]

    def fetch_latest(self) -> list[dict]:
        today = datetime.now(timezone.utc).date()
        records = self.fetch_history(today - timedelta(days=366), today)
        latest = {}
        for record in records:
            if record["fund"] not in latest or record["date"] > latest[record["fund"]]["date"]:
                latest[record["fund"]] = record
        return list(latest.values())
